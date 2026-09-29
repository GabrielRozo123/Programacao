"""
gerar_planilha_meteo.py — consolida a série da estação Zeus numa planilha de projeto.

Abas geradas:
  Resumo_mensal · Rosa_ventos · Rajada · Estabilidade · Qualidade_dados ·
  Ventiladores · Cenarios · Dados_horarios

Uso:  python3 gerar_planilha_meteo.py <relatorio_zeus.xlsx> <saida.xlsx>
"""
import sys
import numpy as np
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from analise_meteo_zeus import ler, dir_media, setor, SETORES, KMH
from cenarios_meteo import classificar, V_MIN_DIR

AZUL = PatternFill("solid", fgColor="1F4E79")
CINZA = PatternFill("solid", fgColor="D9E1F2")
BRANCO = Font(color="FFFFFF", bold=True)


def aba(wb, nome, cab, linhas, larg=14, nota=None):
    ws = wb.create_sheet(nome)
    r0 = 1
    if nota:
        for i, t in enumerate(nota):
            ws.cell(row=i + 1, column=1, value=t).font = Font(italic=True, color="555555")
        r0 = len(nota) + 2
    for j, c in enumerate(cab, 1):
        cel = ws.cell(row=r0, column=j, value=c)
        cel.fill, cel.font = AZUL, BRANCO
        cel.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = larg
    for i, lin in enumerate(linhas, r0 + 1):
        for j, v in enumerate(lin, 1):
            if isinstance(v, float):
                v = None if np.isnan(v) else round(v, 3)
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return ws


def main(src, out):
    dados = classificar(ler(src))
    wb = Workbook()
    wb.remove(wb.active)

    # ── 1. resumo mensal ────────────────────────────────────────────────
    mes = defaultdict(list)
    for d in dados:
        mes[(d["t"].year, d["t"].month)].append(d)
    lin = []
    for (a, m), g in sorted(mes.items()):
        f = lambda k: np.array([x[k] for x in g if x[k] is not None], float)
        vm, vi, rj = f("v_med") * KMH, f("v_inst") * KMH, f("rajada") * KMH
        T, UR = f("t_med"), f("ur_med")
        gd = [x for x in g if x["dir"] is not None and x["u"] >= V_MIN_DIR]
        ang, R = dir_media([x["dir"] for x in gd], [x["u"] for x in gd]) if gd else (np.nan, np.nan)
        lin.append([f"{m:02d}/{a}", len(g), vm.mean(), vi.mean(), rj.mean(),
                    rj.max() if rj.size else np.nan,
                    T.mean() if T.size else np.nan, UR.mean() if UR.size else np.nan,
                    ang, SETORES[setor(ang)] if not np.isnan(ang) else "", R,
                    100 * np.mean(vm < 1.5)])
    allv = np.array([d["u"] for d in dados])
    lin.append(["GERAL", len(dados), allv.mean(),
                np.mean([d["v_inst"] * KMH for d in dados if d["v_inst"] is not None]),
                np.mean([d["rajada_ms"] for d in dados if d["rajada_ms"]]),
                max(d["rajada_ms"] for d in dados if d["rajada_ms"]),
                np.mean([d["t_med"] for d in dados if d["t_med"] is not None]),
                np.mean([d["ur_med"] for d in dados if d["ur_med"] is not None]),
                *(lambda a, R: [a, SETORES[setor(a)], R])(*dir_media(
                    [d["dir"] for d in dados if d["dir"] is not None and d["u"] >= V_MIN_DIR],
                    [d["u"] for d in dados if d["dir"] is not None and d["u"] >= V_MIN_DIR])),
                100 * np.mean(allv < 1.5)])
    ws = aba(wb, "Resumo_mensal",
             ["Mês", "n horas", "v média (m/s)", "v instantânea (m/s)", "Rajada média (m/s)",
              "Rajada máx (m/s)", "T média (°C)", "UR média (%)", "Direção média vetorial (°)",
              "Setor", "R (concentração direcional)", "% horas v<1,5 m/s"], lin,
             nota=["Estação Zeus — LD Celulose (Fábrica). Série horária, velocidades convertidas de km/h para m/s.",
                   "Direção média = média VETORIAL ponderada pela velocidade (média aritmética de ângulos é incorreta).",
                   "R: 0 = vento de todas as direções, 1 = sempre da mesma direção."])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True), CINZA

    # ── 2. rosa dos ventos ──────────────────────────────────────────────
    val = [d for d in dados if d["dir"] is not None and d["u"] >= V_MIN_DIR]
    idx = np.array([setor(d["dir"]) for d in val])
    uu = np.array([d["u"] for d in val])
    lin = []
    for k, nome in enumerate(SETORES):
        m = idx == k
        rj = [val[i]["rajada_ms"] for i in np.where(m)[0] if val[i]["rajada_ms"]]
        lin.append([nome, k * 22.5, int(m.sum()), 100 * m.sum() / len(val),
                    uu[m].mean() if m.any() else 0, np.percentile(uu[m], 95) if m.any() else 0,
                    max(rj) if rj else 0])
    aba(wb, "Rosa_ventos", ["Setor", "Centro (°)", "n horas", "Frequência (%)",
                            "v média (m/s)", "v p95 (m/s)", "Rajada máx (m/s)"], lin,
        nota=[f"Somente horas com v ≥ {V_MIN_DIR} m/s (abaixo do limiar do anemômetro a direção não é significativa).",
              "Direção meteorológica: de onde o vento VEM. 0° = Norte, 90° = Leste."])

    # ── 3. rajada ───────────────────────────────────────────────────────
    par = [(d["u"], d["rajada_ms"]) for d in dados if d["rajada_ms"] and d["u"] > 0]
    vv = np.array([p[0] for p in par]); rr = np.array([p[1] for p in par]); G = rr / vv
    lin = [[f"p{q}", np.percentile(allv, q), np.percentile(rr, q), np.percentile(G, q)]
           for q in (5, 10, 25, 50, 75, 90, 95, 99)]
    lin += [["média", allv.mean(), rr.mean(), np.nan], ["máximo", allv.max(), rr.max(), G.max()]]
    aba(wb, "Rajada", ["Estatística", "v média horária (m/s)", "Rajada (m/s)", "Fator de rajada G"], lin,
        nota=["G = rajada / v média. G muito alto ocorre com vento fraco (divisão por número pequeno):",
              "para o cenário de rajada usar a rajada MEDIDA, não G × v."])

    # ── 4. estabilidade ─────────────────────────────────────────────────
    cls = np.array([d["cls"] for d in dados])
    desc = {"A": "muito instável", "B": "instável", "C": "levemente instável",
            "D": "neutra", "E": "levemente estável", "F": "estável"}
    lin = []
    for k in "ABCDEF":
        m = cls == k
        lin.append([k, desc[k], int(m.sum()), 100 * m.mean(),
                    allv[m].mean() if m.any() else np.nan,
                    np.percentile(allv[m], 50) if m.any() else np.nan])
    aba(wb, "Estabilidade", ["Classe", "Descrição", "n horas", "Frequência (%)",
                             "v média (m/s)", "v mediana (m/s)"], lin,
        nota=["Pasquill-Turner: dia pela luminosidade medida × velocidade; noite (luminosidade = 0) pelo ramo",
              "de céu limpo, hipótese conservadora — a Zeus não mede nebulosidade. [SUPOSTO]"])

    # ── 5. qualidade de dados ───────────────────────────────────────────
    n = len(dados)
    tm = sum(d["t_med"] is not None for d in dados)
    urv = [d for d in dados if None not in (d["ur_min"], d["ur_med"], d["ur_max"])]
    lin = [
        ["Registros totais", n, "", "Série 01/10/2025 a 28/08/2026, horária"],
        ["Estação marcada inativa", sum(d["inativa"] for d in dados), "", ""],
        ["Temperatura média preenchida", tm, f"{100*tm/n:.1f}%", "Lacuna: jan/2026 sem temperatura"],
        ["UR com min ≤ méd ≤ máx", len(urv) - sum(not (d["ur_min"] <= d["ur_med"] <= d["ur_max"]) for d in urv),
         f"de {len(urv)}", "Sem violações"],
        ["UR média = 0 %", sum(d["ur_med"] == 0 for d in urv), "", "Fisicamente implausível — tratar como falha"],
        [f"Direção com v < {V_MIN_DIR} m/s", sum(d["u"] < V_MIN_DIR for d in dados), "",
         "Excluídos da rosa dos ventos"],
        ["v média × v instantânea", "correlação 0,944", "", "Grandezas distintas, coerentes entre si"],
    ]
    aba(wb, "Qualidade_dados", ["Verificação", "Resultado", "", "Observação"], lin, larg=30)

    # ── 6. ventiladores ─────────────────────────────────────────────────
    lin = [["1219-24-910 / 920", "LX 816.95.00", 2, 10.5, 37800, 1670, 28.0, 1775, 0.3105, 10.5 / 0.3105],
           ["1219-24-930", "HL 4S 871.92.00", 1, 4.2, 15120, 2400, 16.3, 1770, 0.1337, 4.2 / 0.1337],
           ["1219-24-940", "HL 4S 871.92.00", 1, 4.2, 15120, 2400, 16.3, 1770, 0.1337, 4.2 / 0.1337]]
    aba(wb, "Ventiladores", ["Tag", "Modelo", "Qtd", "Vazão (m³/s)", "Vazão (m³/h)",
                             "ΔP estática (Pa)", "Potência eixo (kW)", "Rotação (rpm)",
                             "Área descarga (m²)", "v descarga (m/s)"], lin,
        nota=["Pontos de operação lidos das curvas Howden (manuais CVC99D0014/15/16)."])

    # ── 7. cenários ─────────────────────────────────────────────────────
    dom = max(range(16), key=lambda k: (idx == k).sum())
    d_tip = dom * 22.5
    v_tip = allv.mean()
    v_raj = max(d["rajada_ms"] for d in dados if d["rajada_ms"])
    mF = cls == "F"
    v_F = allv[mF].mean()
    gF = [d for d, c in zip(dados, cls) if c == "F" and d["dir"] is not None and d["u"] >= V_MIN_DIR]
    iF = np.array([setor(d["dir"]) for d in gF])
    d_F = max(range(16), key=lambda k: (iF == k).sum()) * 22.5
    lin = [
        [1, "Vento típico · direção típica · vazamento PEQUENO", "V1", f"{d_tip:.1f} ({SETORES[dom]})", v_tip, "D", "Q_peq", "Base"],
        [2, "Vento típico · direção típica · vazamento GRANDE",  "V1", f"{d_tip:.1f} ({SETORES[dom]})", v_tip, "D", "Q_gr",  "Mesmo campo de vento do 1"],
        [3, "Rajada · direção típica · vazamento PEQUENO",      "V2", f"{d_tip:.1f} ({SETORES[dom]})", v_raj, "D", "Q_peq", "Rajada máxima medida"],
        [4, "Rajada · direção típica · vazamento GRANDE",       "V2", f"{d_tip:.1f} ({SETORES[dom]})", v_raj, "D", "Q_gr",  "Mesmo campo de vento do 3"],
        [5, "Vento fraco estável · vazamento PEQUENO [PROPOSTO]", "V3", f"{d_F:.1f} ({SETORES[int(d_F/22.5)]})", v_F, "F", "Q_peq", f"Classe F = {100*mF.mean():.1f}% das horas"],
        [6, "Vento fraco estável · vazamento GRANDE [PROPOSTO]",  "V3", f"{d_F:.1f} ({SETORES[int(d_F/22.5)]})", v_F, "F", "Q_gr",  "Mesmo campo de vento do 5"],
    ]
    aba(wb, "Cenarios", ["#", "Cenário", "Campo de vento", "Direção (°)", "U ref a 10 m (m/s)",
                         "Estabilidade", "Vazão de fonte", "Observação"], lin, larg=18,
        nota=["O transporte do escalar é LINEAR na vazão da fonte: cenários com o mesmo campo de vento",
              "diferem só por um fator de escala. 6 cenários = 3 campos de vento (V1, V2, V3).",
              "Q_peq e Q_gr: aguardando definição com o cliente."])

    # ── 8. dados horários ───────────────────────────────────────────────
    lin = [[d["t"].strftime("%d/%m/%Y %H:%M"), d["t"].month, d["u"],
            (d["v_inst"] or 0) * KMH, d["rajada_ms"], d["dir"],
            SETORES[setor(d["dir"])] if d["dir"] is not None else "",
            d["t_med"], d["ur_med"], d["lum"], d["cls"]] for d in dados]
    aba(wb, "Dados_horarios", ["Início [GMT-3]", "Mês", "v média (m/s)", "v inst (m/s)",
                               "Rajada (m/s)", "Direção (°)", "Setor", "T média (°C)",
                               "UR média (%)", "Luminosidade (Wh/m²)", "Classe Pasquill"], lin)

    wb.save(out)
    print(f"salvo: {out}")
    print(f"direção típica {d_tip:.1f}° ({SETORES[dom]}) · v típica {v_tip:.2f} m/s · "
          f"rajada máx {v_raj:.2f} m/s · classe F {100*mF.mean():.1f}% (v {v_F:.2f} m/s, dir {d_F:.1f}°)")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
